import os
import zipfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Define output directory
output_dir = r"c:\Users\admin\Documents\VS Tím\web_homestay\web_homestay\docs\demo-import-word"
assets_dir = os.path.join(output_dir, "assets-room-import")

# 1. Branch Template Data
branch_headers = ["Name", "Address", "Description", "Hotline", "Email", "MapUrl", "BookingLeadTimeHours"]
branch_data = [
    [
        "LumiStay Phú Nhuận",
        "18 Hoa Sứ, Phú Nhuận, TP.HCM",
        "Chi nhánh gần sân bay, tiện cho khách công tác.",
        "0901000011",
        "phunhuan@lumistay.vn",
        "https://maps.google.com/?q=18+Hoa+Su+Phu+Nhuan",
        2
    ],
    [
        "LumiStay Gò Vấp",
        "125 Quang Trung, Gò Vấp, TP.HCM",
        "Phù hợp khách đi ngắn ngày và cặp đôi.",
        "0901000012",
        "govap@lumistay.vn",
        "https://maps.google.com/?q=125+Quang+Trung+Go+Vap",
        3
    ],
    [
        "LumiStay Thủ Đức",
        "9 Đặng Văn Bi, TP. Thủ Đức",
        "Không gian yên tĩnh cho khách ở linh hoạt theo giờ.",
        "0901000013",
        "thuduc@lumistay.vn",
        "https://maps.google.com/?q=9+Dang+Van+Bi+Thu+Duc",
        2
    ]
]

# 2. Room Template Data
room_headers = [
    "Name", "Description", "PricePerHour", "PricePerDay", "ExtraGuestFee",
    "PriceWeekendPerHour", "PriceWeekendPerDay", "PriceHolidayPerHour", "PriceHolidayPerDay",
    "Capacity", "MaxGuests", "Status", "BranchName", "AvatarFile", "AdditionalFiles"
]
room_data = [
    [
        "Lụa Mây PN-201",
        "Phòng sáng, nhẹ nhàng, phù hợp cặp đôi hoặc khách công tác.",
        160000, 1100000, 100000, 185000, 1260000, 210000, 1450000,
        2, 4, "Available", "LumiStay Phú Nhuận", "room-main-01.jpg", "room-extra-01.jpg, room-extra-02.jpg"
    ],
    [
        "Hoàng Kim GV-501",
        "Phòng rộng, cao cấp, hợp khách nghỉ dưỡng ngắn ngày.",
        280000, 2100000, 150000, 320000, 2450000, 360000, 2780000,
        4, 6, "Available", "LumiStay Gò Vấp", "room-main-02.jpg", "room-extra-03.jpg, room-extra-04.jpg"
    ],
    [
        "An Nhiên TD-101",
        "Phòng tiêu chuẩn, sạch đẹp, tối ưu đặt nhanh theo giờ.",
        120000, 780000, 80000, 140000, 900000, 155000, 990000,
        2, 3, "Available", "LumiStay Thủ Đức", "room-main-03.png", "room-extra-05.jpg"
    ]
]

# 3. Staff Template Data
staff_headers = ["Username", "FullName", "Role", "BranchName", "Permissions"]
staff_data = [
    [
        "pn.manager",
        "Nguyễn Minh Phúc",
        "Manager",
        "LumiStay Phú Nhuận",
        "matrix.view, bookings.view, bookings.detail, rooms.view, chat.view, chat.reply"
    ],
    [
        "gv.staff01",
        "Trần Mỹ Linh",
        "Staff",
        "LumiStay Gò Vấp",
        "matrix.view, bookings.view, bookings.detail, chat.view, chat.reply"
    ],
    [
        "thuduc.admin",
        "Phạm Quốc An",
        "Manager",
        "LumiStay Thủ Đức",
        "matrix.view, bookings.view, branches.view, rooms.view, statistics.view, chat.view"
    ]
]

# 4. AI Knowledge Template Data
ai_headers = ["ScopeName", "Title", "Content", "Tags", "Priority", "IsActive"]
ai_data = [
    [
        "Public Booking",
        "Đặt phòng theo giờ",
        "Khách có thể đặt theo giờ nếu phòng còn slot hợp lệ theo dữ liệu thật của room_slot_inventories.",
        "dat-phong, theo-gio, slot",
        5,
        "Yes"
    ],
    [
        "Cancellation",
        "Yêu cầu hủy cần ảnh xác nhận",
        "Khi gửi yêu cầu hủy, khách cần cung cấp ảnh xác nhận email hoặc bằng chứng hợp lệ để nhân viên duyệt nhanh.",
        "huy-don, refund, xac-nhan",
        4,
        "Yes"
    ],
    [
        "Operations",
        "Lead time chi nhánh",
        "Mỗi chi nhánh có thể có lead time riêng, nếu không có sẽ dùng lead time toàn hệ thống.",
        "lead-time, chi-nhanh",
        3,
        "Yes"
    ]
]

# 5. Slot Template Data
slot_headers = [
    "Name", "Code", "DurationMinutes", "CleanupMinutes", "SeedStartTime",
    "FixedStartTime", "FixedEndTime", "CrossesMidnight", "IsActive"
]
slot_data = [
    [
        "Combo 6 tiếng demo",
        "combo_6h_demo",
        360,
        30,
        "08:00",
        "",
        "",
        "No",
        "Yes"
    ],
    [
        "Combo 8 tiếng demo",
        "combo_8h_demo",
        480,
        30,
        "08:00",
        "",
        "",
        "No",
        "Yes"
    ],
    [
        "Ca đêm cố định",
        "fixed_night_demo",
        480,
        15,
        "",
        "22:00",
        "06:00",
        "Yes",
        "Yes"
    ]
]

# 6. Holiday Template Data
holiday_headers = ["Date", "Description"]
holiday_data = [
    ["2026-09-02", "Quốc khánh 2/9"],
    ["2026-12-31", "Đêm cuối năm - áp giá cao điểm"],
    ["2027-01-01", "Tết Dương Lịch"]
]

# Styles
font_header = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
font_data = Font(name="Segoe UI", size=11, color="333333")
fill_header = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid") # Classic Navy Blue
fill_zebra = PatternFill(start_color="F2F6F9", end_color="F2F6F9", fill_type="solid") # Soft blue-gray
border_thin = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='thin', color='D3D3D3')
)
align_left = Alignment(horizontal="left", vertical="center")
align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

def create_excel_file(filename, headers, data_rows, numeric_cols=set(), date_cols=set(), center_cols=set()):
    wb = Workbook()
    ws = wb.active
    ws.title = "Import Template"
    
    # Write headers
    ws.append(headers)
    
    # Write data
    for row in data_rows:
        ws.append(row)
    
    # Set header row height
    ws.row_dimensions[1].height = 28
    
    # Format header cells
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_thin
    
    # Format data cells
    for row_idx in range(2, len(data_rows) + 2):
        ws.row_dimensions[row_idx].height = 22
        is_even = (row_idx % 2 == 0)
        
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = font_data
            cell.border = border_thin
            
            # Zebra striping
            if is_even:
                cell.fill = fill_zebra
            
            # Alignments & Number formatting
            if col_idx in numeric_cols:
                cell.alignment = align_right
                cell.number_format = '#,##0'
            elif col_idx in date_cols:
                cell.alignment = align_center
                cell.number_format = 'yyyy-mm-dd'
            elif col_idx in center_cols:
                cell.alignment = align_center
            else:
                cell.alignment = align_left
                
    # Auto-adjust column widths with safety margin
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            # If formatted numeric, string length is smaller than actual display, so we add a buffer
            val_len = len(val_str)
            if cell.column in numeric_cols and cell.row > 1:
                val_len += 3 # padding for commas
            if val_len > max_len:
                max_len = val_len
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    # Save the file
    file_path = os.path.join(output_dir, filename)
    wb.save(file_path)
    print(f"Created Excel file: {file_path}")
    return file_path

# Generate Excel files
create_excel_file(
    "01-mau-chi-nhanh.xlsx",
    branch_headers,
    branch_data,
    numeric_cols={7}, # BookingLeadTimeHours
    center_cols={4, 7} # Hotline, BookingLeadTimeHours
)

# For room, columns are:
# 1 Name, 2 Desc, 3 PricePerHour, 4 PricePerDay, 5 ExtraGuestFee,
# 6 PriceWeekendPerHour, 7 PriceWeekendPerDay, 8 PriceHolidayPerHour, 9 PriceHolidayPerDay,
# 10 Capacity, 11 MaxGuests, 12 Status, 13 BranchName, 14 AvatarFile, 15 AdditionalFiles
room_excel_path = create_excel_file(
    "02-mau-phong-ngu.xlsx",
    room_headers,
    room_data,
    numeric_cols={3, 4, 5, 6, 7, 8, 9, 10, 11},
    center_cols={10, 11, 12} # Capacity, MaxGuests, Status
)

create_excel_file(
    "03-mau-nhan-vien.xlsx",
    staff_headers,
    staff_data,
    center_cols={1, 3} # Username, Role
)

create_excel_file(
    "04-ngan-hang-tri-thuc-ai.xlsx",
    ai_headers,
    ai_data,
    numeric_cols={5}, # Priority
    center_cols={5, 6} # Priority, IsActive
)

# 1 Name, 2 Code, 3 DurationMinutes, 4 CleanupMinutes, 5 SeedStartTime,
# 6 FixedStartTime, 7 FixedEndTime, 8 CrossesMidnight, 9 IsActive
create_excel_file(
    "05-mau-khung-gio.xlsx",
    slot_headers,
    slot_data,
    numeric_cols={3, 4}, # DurationMinutes, CleanupMinutes
    center_cols={2, 3, 4, 5, 6, 7, 8, 9} # Code, Duration, Cleanup, StartTimes, CrossesMidnight, IsActive
)

create_excel_file(
    "06-mau-ngay-le.xlsx",
    holiday_headers,
    holiday_data,
    date_cols={1} # Date
)

# 7. Package ZIP file 02-mau-phong-ngu.zip
# According to instructions, this ZIP must contain:
# 1. 02-mau-phong-ngu.xlsx at the root
# 2. The images from assets-room-import/ directory at the root of the ZIP
zip_path = os.path.join(output_dir, "02-mau-phong-ngu.zip")
with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
    # Add room Excel file
    zipf.write(room_excel_path, "02-mau-phong-ngu.xlsx")
    
    # Add images from assets-room-import directory
    if os.path.exists(assets_dir):
        for img_file in os.listdir(assets_dir):
            img_path = os.path.join(assets_dir, img_file)
            if os.path.isfile(img_path) and img_file.lower().endswith(('.jpg', '.jpeg', '.png', '.webp', '.gif')):
                # Write with arcname at the root of the zip
                zipf.write(img_path, img_file)
                print(f"Added to ZIP: {img_file}")
                
print(f"Created ZIP file: {zip_path}")

# Clean up temporary 02-mau-phong-ngu.xlsx at output_dir
# Wait, should we keep it? The user said "5 file excel và 1 file zip cho 6 mẫu import". 
# So there are 6 templates in total: 
# - Chi nhánh (01), Nhân viên (03), Tri thức (04), Khung giờ (05), Ngày lễ (06) are the 5 Excel files.
# - Phòng ngủ (02) is the 1 ZIP file.
# If we keep the 02-mau-phong-ngu.xlsx outside the ZIP, it would be a 6th Excel file.
# The user wants exactly "5 file excel và 1 file zip cho 6 mẫu import".
# So yes, 02-mau-phong-ngu.xlsx should ONLY be inside the 02-mau-phong-ngu.zip!
# We will delete the temporary 02-mau-phong-ngu.xlsx to keep the output directory clean and match the "5 file excel và 1 file zip" count.
if os.path.exists(room_excel_path):
    os.remove(room_excel_path)
    print(f"Removed temporary Excel file from output directory: {room_excel_path}")
